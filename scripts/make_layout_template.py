#!/usr/bin/env python3
"""Generate a blank layout-only workbook: Plate Layout + Standards + Samples,
with no Raw Data sheet. Pair this with a native instrument export (e.g. a
Tecan Spark / SparkControl "Result sheet" .xlsx) via:

    python -m elisa_analysis.cli --input my_tecan_export.xlsx \\
        --layout my_kit_layout.xlsx

Since the layout (which wells are standards/samples, concentrations,
dilutions) is usually reused across many runs of the same ELISA kit, you
typically only need to fill this in once.

Usage:
    python scripts/make_layout_template.py [output_path]
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROWS = list("ABCDEFGH")
COLS = list(range(1, 13))
HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
HEADER_FONT = Font(bold=True)


def _write_grid(ws, corner_text, fill_value=None):
    ws.cell(row=1, column=1, value=corner_text).font = HEADER_FONT
    for j, col_num in enumerate(COLS):
        c = ws.cell(row=1, column=2 + j, value=col_num)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
    for i, row_letter in enumerate(ROWS):
        c = ws.cell(row=2 + i, column=1, value=row_letter)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        for j, col_num in enumerate(COLS):
            cell = ws.cell(row=2 + i, column=2 + j)
            if fill_value is not None:
                cell.value = fill_value(row_letter, col_num)
    ws.column_dimensions["A"].width = 8
    for j in range(len(COLS)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 10


def build_layout_template() -> Workbook:
    wb = Workbook()

    ws_layout = wb.active
    ws_layout.title = "Plate Layout"

    def default_layout(row_letter, col_num):
        if col_num in (1, 2):
            idx = ROWS.index(row_letter) + 1
            return f"STD{idx}"
        return ""

    _write_grid(ws_layout, "Well", fill_value=default_layout)
    ws_layout["A11"] = (
        "Label each well to match the physical layout of your plate: STDn for "
        "standards (must match the Standards sheet), BLANK for zero/blank wells, "
        "and any short code (e.g. UNK1) for sample wells. Repeat a label across "
        "wells for replicates. This must line up with the well positions actually "
        "read by the instrument in your raw-data export."
    )
    ws_layout["A11"].font = Font(italic=True, size=9)

    ws_std = wb.create_sheet("Standards")
    ws_std.append(["Label", "Concentration"])
    for c in ws_std[1]:
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i in range(8, 0, -1):
        ws_std.append([f"STD{i}", ""])
    ws_std["D1"] = "Enter the known concentration for each standard label used in Plate Layout."
    ws_std["D1"].font = Font(italic=True, size=9)
    ws_std.column_dimensions["A"].width = 12
    ws_std.column_dimensions["B"].width = 14
    ws_std.column_dimensions["D"].width = 70

    ws_smp = wb.create_sheet("Samples")
    ws_smp.append(["Label", "Sample Name", "Dilution Factor"])
    for c in ws_smp[1]:
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    ws_smp.append(["UNK1", "Patient 001", 1])
    ws_smp["E1"] = (
        "Optional: map well labels (e.g. UNK1) to a readable sample name and a dilution "
        "factor applied after interpolation. Omit a sample here to default to Dilution=1."
    )
    ws_smp["E1"].font = Font(italic=True, size=9)
    ws_smp.column_dimensions["A"].width = 12
    ws_smp.column_dimensions["B"].width = 20
    ws_smp.column_dimensions["C"].width = 16
    ws_smp.column_dimensions["E"].width = 90

    return wb


if __name__ == "__main__":
    out_path = sys.argv[1] if len(sys.argv) > 1 else "elisa_layout_template.xlsx"
    build_layout_template().save(out_path)
    print(f"Wrote layout template to {out_path}")
