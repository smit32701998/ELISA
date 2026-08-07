#!/usr/bin/env python3
"""Generate a synthetic Tecan Spark / SparkControl-style "Result sheet" export,
plus a matching layout workbook, to demo/test elisa_analysis against a native
instrument export rather than the combined template.

Mimics the real SparkControl export structure: a metadata header block,
followed by three 8x12 grids (raw absorbance, Reference, Difference) each
introduced by a "<>" corner cell -- the Difference table (reference-wavelength
corrected) is what elisa_analysis will auto-select for analysis.

Usage:
    python scripts/make_tecan_example.py [raw_output_path] [layout_output_path]
"""
import sys

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, ".")
from elisa_analysis.fitting import logistic_4pl

ROWS = list("ABCDEFGH")
COLS = list(range(1, 13))
TRUE_PARAMS = dict(a=0.05, b=1.15, c=25.0, d=2.9)
STD_CONCS = [100, 50, 25, 12.5, 6.25, 3.125, 1.5625, 0]
REFERENCE_OD = 0.045  # roughly constant reference-wavelength background


def _write_metadata_header(ws, assay_name="Synthetic IL6 Demo"):
    rows = [
        ["Method name: Method 1"],
        ["Application: SparkControl", None, None, None, "V3.1 SP1"],
        ["Device: Spark", None, None, None, "Serial number: 0000000000"],
        ["Firmware:", None, None, None, "ABS:V4.3.2"],
        [""],
        ["Date:", None, None, None, "2026-08-07"],
        ["Time:", None, None, None, "10:00 AM"],
        ["System", None, None, None, "DEMO-PC"],
        ["User", None, None, None, "DEMO-PC\\demo"],
        ["Plate", None, None, None, "[NUN96ft] - generic 96-well flat-bottom plate"],
        ["Lid lifter", None, None, None, "No lid"],
        ["Humidity Cassette", None, None, None, "No humidity cassette"],
        ["Smooth mode", None, None, None, "Not selected"],
        [],
        ["List of actions in this measurement script:"],
        ["Plate"],
        [None, "Absorbance", None, None, None, None, assay_name],
        [],
        ["Name", None, None, None, "NUN96ft"],
        ["Plate layout"],
        ["Plate area", None, None, None, "A1-H12"],
        [],
        ["Mode", "Absorbance"],
        ["Name", assay_name],
        ["Measurement wavelength [nm]", None, None, None, 450],
        ["Reference wavelength [nm]", None, None, None, 570],
        ["Number of flashes", None, None, None, 10],
        ["Settle time [ms]", None, None, None, 50],
        ["Part of Plate", None, None, None, "A1-H12"],
        [],
        ["Start Time", None, None, None, "2026-08-07 10:00:00"],
        ["Temperature [°C]", None, None, None, 37.0],
        [],
        [assay_name],
    ]
    for row in rows:
        ws.append(row)


def _write_grid_table(ws, label, values_by_well, include_label=True):
    if include_label:
        ws.append([label])
    ws.append(["<>"] + [str(c) for c in COLS])
    for row_letter in ROWS:
        ws.append([row_letter] + [values_by_well.get(f"{row_letter}{c}", "") for c in COLS])
    ws.append([])
    ws.append([])


def build_tecan_example(seed: int = 7):
    rng = np.random.default_rng(seed)

    raw = {}
    for i, row_letter in enumerate(ROWS):
        label = f"STD{i + 1}"
        conc = STD_CONCS[i]
        for col_num in (1, 2):
            true_y = logistic_4pl(np.array([conc if conc > 0 else 1e-6]), **TRUE_PARAMS)[0]
            if conc == 0:
                true_y = TRUE_PARAMS["a"]
            noisy = true_y * (1 + rng.normal(0, 0.05)) + rng.normal(0, 0.004)
            raw[f"{row_letter}{col_num}"] = round(max(noisy, 0.0) + REFERENCE_OD, 4)

    sample_concs = {"UNK1": 40.0, "UNK2": 18.0, "UNK3": 8.0, "UNK4": 60.0, "UNK5": 2.0}
    wells_cycle = [(r, c) for c in range(3, 13) for r in ROWS]
    wi = 0
    sample_layout = {}
    for name, conc in sample_concs.items():
        true_y = logistic_4pl(np.array([conc]), **TRUE_PARAMS)[0]
        for _rep in range(2):
            row_letter, col_num = wells_cycle[wi]
            wi += 1
            noisy = true_y * (1 + rng.normal(0, 0.06)) + rng.normal(0, 0.004)
            raw[f"{row_letter}{col_num}"] = round(max(noisy, 0.0) + REFERENCE_OD, 4)
            sample_layout[f"{row_letter}{col_num}"] = name

    reference = {f"{r}{c}": round(REFERENCE_OD + rng.normal(0, 0.003), 4) for r in ROWS for c in COLS}
    difference = {
        wid: round(raw.get(wid, 0) - reference.get(wid, 0), 4) for wid in raw
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Result sheet"
    _write_metadata_header(ws)
    _write_grid_table(ws, "Synthetic IL6 Demo", raw, include_label=False)
    _write_grid_table(ws, "Reference", reference)
    _write_grid_table(ws, "Difference", difference)

    layout_wb = Workbook()
    ws_layout = layout_wb.active
    ws_layout.title = "Plate Layout"
    ws_layout.append(["Well"] + [str(c) for c in COLS])
    for row_letter in ROWS:
        row = [row_letter]
        for col_num in COLS:
            wid = f"{row_letter}{col_num}"
            if col_num in (1, 2):
                idx = ROWS.index(row_letter) + 1
                row.append(f"STD{idx}")
            else:
                row.append(sample_layout.get(wid, ""))
        ws_layout.append(row)

    ws_std = layout_wb.create_sheet("Standards")
    ws_std.append(["Label", "Concentration"])
    for i, conc in enumerate(STD_CONCS):
        ws_std.append([f"STD{i + 1}", conc])

    ws_smp = layout_wb.create_sheet("Samples")
    ws_smp.append(["Label", "Sample Name", "Dilution Factor"])
    for name in sample_concs:
        ws_smp.append([name, name.replace("UNK", "Patient "), 1])

    return wb, layout_wb


if __name__ == "__main__":
    raw_out = sys.argv[1] if len(sys.argv) > 1 else "examples/tecan_spark_raw_example.xlsx"
    layout_out = sys.argv[2] if len(sys.argv) > 2 else "examples/tecan_spark_layout_example.xlsx"
    raw_wb, layout_wb = build_tecan_example()
    raw_wb.save(raw_out)
    layout_wb.save(layout_out)
    print(f"Wrote {raw_out} and {layout_out}")
