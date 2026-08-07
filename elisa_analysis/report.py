"""Write a formatted Excel report (results tables + embedded Prism-style charts)."""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from .analysis import AnalysisResult
from .plotting import save_all_plots

HEADER_FILL = PatternFill(start_color="2E3192", end_color="2E3192", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FLAG_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")


def _write_df(ws, df, start_row=1, number_formats=None):
    number_formats = number_formats or {}
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start_row):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == start_row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center")
            else:
                col_name = df.columns[c_idx - 1]
                if col_name in number_formats:
                    cell.number_format = number_formats[col_name]
    for c_idx, col in enumerate(df.columns, start=1):
        max_len = df[col].astype(str).str.len().max()
        max_len = 10 if pd.isna(max_len) else max_len
        width = max(12, min(28, int(max_len) + 4))
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    return start_row + len(df) + 1


def write_report(result: AnalysisResult, out_path: str, plot_prefix: str) -> str:
    plot_paths = save_all_plots(result, plot_prefix)

    wb = Workbook()

    ws_sum = wb.active
    ws_sum.title = "Summary"
    report_title = f"{result.analyte} ELISA Analysis Report" if result.analyte else "ELISA Analysis Report"
    ws_sum["A1"] = report_title
    ws_sum["A1"].font = Font(size=16, bold=True)
    wb.properties.title = report_title
    ws_sum["A3"] = "Curve model"
    ws_sum["B3"] = result.fit.model
    ws_sum["A4"] = "R²"
    ws_sum["B4"] = round(result.fit.r_squared, 5)
    ws_sum["A5"] = "Blank OD subtracted"
    ws_sum["B5"] = round(result.blank_od, 5)
    row = 7
    if result.available_tables and len(result.available_tables) > 1:
        ws_sum["A6"] = "OD table used"
        ws_sum["B6"] = result.source_table
        ws_sum["A7"] = "Tables found in source"
        ws_sum["B7"] = ", ".join(result.available_tables)
        row = 9
    for name, value in result.fit.params.items():
        ws_sum.cell(row=row, column=1, value=name)
        ws_sum.cell(row=row, column=2, value=round(float(value), 6))
        row += 1
    if result.fit.warnings:
        row += 1
        ws_sum.cell(row=row, column=1, value="FIT WARNINGS").font = Font(color="CC0000", bold=True)
        row += 1
        for w in result.fit.warnings:
            ws_sum.cell(row=row, column=1, value=w).font = Font(color="CC0000")
            ws_sum.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            ws_sum.row_dimensions[row].height = 45
            row += 1
    ws_sum.column_dimensions["A"].width = 22
    ws_sum.column_dimensions["B"].width = 18

    img1 = XLImage(plot_paths[0])
    img1.width, img1.height = 520, 400
    ws_sum.add_image(img1, "D3")

    img2 = XLImage(plot_paths[1])
    img2.width, img2.height = 520, 400
    ws_sum.add_image(img2, "D25")

    ws_std = wb.create_sheet("Standards")
    std_out = result.standards_table[
        ["Label", "Concentration", "N", "MeanOD", "SD_OD", "PctCV", "CorrectedOD", "BackCalcConc", "PctRecovery", "Flag"]
    ].copy()
    end_row = _write_df(
        ws_std,
        std_out,
        number_formats={
            "MeanOD": "0.0000",
            "SD_OD": "0.0000",
            "PctCV": "0.0",
            "CorrectedOD": "0.0000",
            "BackCalcConc": "0.000",
            "PctRecovery": "0.0",
        },
    )
    for r in range(2, end_row):
        if ws_std.cell(row=r, column=std_out.columns.get_loc("Flag") + 1).value:
            for c in range(1, len(std_out.columns) + 1):
                ws_std.cell(row=r, column=c).fill = FLAG_FILL

    ws_smp = wb.create_sheet("Samples")
    if not result.samples_table.empty:
        smp_out = result.samples_table[
            ["Label", "SampleName", "N", "MeanOD", "SD_OD", "PctCV", "CorrectedOD", "Dilution", "InterpolatedConc", "FinalConc", "Flag"]
        ].copy()
        end_row = _write_df(
            ws_smp,
            smp_out,
            number_formats={
                "MeanOD": "0.0000",
                "SD_OD": "0.0000",
                "PctCV": "0.0",
                "CorrectedOD": "0.0000",
                "InterpolatedConc": "0.000",
                "FinalConc": "0.000",
            },
        )
        for r in range(2, end_row):
            if ws_smp.cell(row=r, column=smp_out.columns.get_loc("Flag") + 1).value:
                for c in range(1, len(smp_out.columns) + 1):
                    ws_smp.cell(row=r, column=c).fill = FLAG_FILL
    else:
        ws_smp["A1"] = "No sample wells were found in the Plate Layout."

    ws_raw = wb.create_sheet("Raw Wells")
    _write_df(ws_raw, result.plate_df)

    wb.save(out_path)
    return out_path
